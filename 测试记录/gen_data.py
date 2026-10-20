# -*- coding: utf-8 -*-
"""生成文件格式矩阵所需的测试数据，输出到 测试记录/测试数据/。"""
import os
import shutil
import sys

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "测试数据")
os.makedirs(OUT, exist_ok=True)

MANUAL = os.path.join(BASE, "..", "data", "产品手册.txt")
manual_text = open(MANUAL, encoding="utf-8").read()


def p(name):
    return os.path.join(OUT, name)


# 1. markdown
md = """# 云帆 X3 智能旅行规划仪 产品手册（V2.3）

## 一、产品基本信息
- 产品型号：YF-X3-2026
- 官方售价：标准版 1899 元，尊享版 2499 元
- 机身重量：238 克
- 电池容量：4200 毫安时

## 二、核心功能
1. 行程自动规划
2. 离线地图（全国 342 个城市）
3. 实时翻译（42 种语言）

> 注：本文件用于测试 markdown 解析能力。
"""
open(p("产品手册.md"), "w", encoding="utf-8").write(md)

# 2. csv
csv = """型号,版本,官方售价(元),重量(克),电池容量(毫安时),防水等级
YF-X3-2026,标准版,1899,238,4200,IPX7
YF-X3-2026,尊享版,2499,238,4200,IPX7
YF-X2-2025,标准版,1499,245,3800,IPX6
"""
open(p("产品价格表.csv"), "w", encoding="utf-8-sig").write(csv)

# 3. xlsx（两个 sheet）
import openpyxl
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "价格表"
for row in [
    ["型号", "版本", "官方售价(元)", "电池容量(毫安时)"],
    ["YF-X3-2026", "标准版", 1899, 4200],
    ["YF-X3-2026", "尊享版", 2499, 4200],
]:
    ws.append(row)
ws2 = wb.create_sheet("售后政策")
for row in [
    ["项目", "内容"],
    ["整机保修", "12 个月"],
    ["电池保修", "6 个月"],
]:
    ws2.append(row)
wb.save(p("产品配置表.xlsx"))

# 4. docx（标题层级 + 表格）
import docx
d = docx.Document()
d.add_heading("云帆 X3 智能旅行规划仪 产品说明书", level=1)
d.add_heading("一、产品基本信息", level=2)
d.add_paragraph("产品型号：YF-X3-2026")
d.add_paragraph("官方售价：标准版 1899 元，尊享版 2499 元")
d.add_paragraph("机身重量：238 克，电池容量：4200 毫安时")
d.add_heading("二、保修政策", level=2)
t = d.add_table(rows=3, cols=2)
t.cell(0, 0).text = "项目"
t.cell(0, 1).text = "内容"
t.cell(1, 0).text = "整机保修"
t.cell(1, 1).text = "12 个月"
t.cell(2, 0).text = "电池保修"
t.cell(2, 1).text = "6 个月"
d.save(p("产品说明书.docx"))

# 5. pdf（复用本机已有中文 PDF）
src_pdf = r"C:\Users\曾\Desktop\云计算导论\实验二\云计算导论 -- 实验二.pdf"
if os.path.exists(src_pdf):
    shutil.copy(src_pdf, p("测试文档.pdf"))
else:
    print("!! 未找到源 PDF，跳过")

# 6. json
import json
data = {
    "product": "云帆 X3 智能旅行规划仪",
    "model": "YF-X3-2026",
    "price": {"标准版": 1899, "尊享版": 2499},
    "spec": {"重量": "238克", "电池容量": "4200毫安时", "防水等级": "IPX7"},
}
open(p("结构化数据.json"), "w", encoding="utf-8").write(
    json.dumps(data, ensure_ascii=False, indent=2))

# 7. html
html = """<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>云帆 X3 产品页</title></head>
<body>
<h1>云帆 X3 智能旅行规划仪</h1>
<p>官方售价：标准版 1899 元，尊享版 2499 元。</p>
<p>电池容量 4200 毫安时，续航 16 小时。</p>
</body></html>
"""
open(p("产品页.html"), "w", encoding="utf-8").write(html)

# 8. 伪装成 exe 的文本文件
open(p("伪装文件.exe"), "wb").write(("MZ" + manual_text).encode("utf-8"))

# 9. 0 字节文件
open(p("空文件.txt"), "wb").close()

# 10. 无扩展名文件
open(p("无扩展名文件"), "w", encoding="utf-8").write(manual_text)

# 11. 超大文件（约 1.5MB）
chunk = manual_text * 1
big = (manual_text + "\n\n") * 900
open(p("大文件.txt"), "w", encoding="utf-8").write(big)

# 12. 补充说明（用于多文档/跨文档/删除测试）
extra = """云帆 X3 企业采购与售后服务补充说明

一、企业采购政策
1. 单笔采购满 50 台可享 8 折优惠，满 100 台可享 7.5 折。
2. 企业客户支持开具增值税专用发票。
3. 企业采购需签订框架协议，账期为 30 天。

二、官方配件价格
1. 旅行收纳包：129 元。
2. 车载支架：89 元。
3. 快充充电器（9V/2A）：99 元。

三、售后服务网点
1. 全国 126 个城市设有授权维修点。
2. 寄修地址：广东省深圳市南山区云帆科技园 A 栋 3 层。
3. 以旧换新：旧机最高可抵扣 300 元。

四、特别说明
本补充说明与产品手册如有冲突，以产品手册为准。
"""
open(p("补充说明.txt"), "w", encoding="utf-8").write(extra)

# 13. 特殊文件名
open(p("产品 手册(最终版)#1.txt"), "w", encoding="utf-8").write(manual_text)

print("生成完成，目录:", OUT)
for f in sorted(os.listdir(OUT)):
    print(f"  {f:32s} {os.path.getsize(os.path.join(OUT, f)):>10,} 字节")
