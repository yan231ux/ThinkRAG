# -*- coding: utf-8 -*-
"""根据 cases_thinkrag.py / bugs_thinkrag.py 生成两份 Excel 报告。

用法：python build_reports.py
新增用例或缺陷后重新运行即可，表格会整体重建（数据源始终以 py 文件为准）。
"""

import os
import sys
from collections import Counter, OrderedDict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from cases_thinkrag import CASES
from bugs_thinkrag import BUGS, CLOSURE

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# ---------- 样式 ----------
FONT_NAME = "微软雅黑"
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
TITLE_FONT = Font(name=FONT_NAME, size=13, bold=True)
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT_NAME, size=9)
BOLD_BODY = Font(name=FONT_NAME, size=9, bold=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

RESULT_COLORS = {
    "通过": "1E7B34",
    "不通过": "C00000",
    "待执行": "808080",
    "阻塞": "BF8F00",
}
PRI_COLORS = {"P0": "C00000", "P1": "C55A11", "P2": "404040", "P3": "808080"}
STATUS_COLORS = {
    "待修复": "C00000",
    "已定位": "C55A11",
    "待复现": "BF8F00",
    "已验证修复": "1E7B34",
}


def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER
    ws.row_dimensions[row].height = 22


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def auto_row_height(ws, row, widths, min_h=18, max_h=260):
    """按单元格内容估算行高，避免打开 Excel 时文字被截断。"""
    import math
    lines = 1
    for i, w in enumerate(widths, start=1):
        v = ws.cell(row=row, column=i).value
        if v is None:
            continue
        text = str(v)
        # 半角字符约占 1 个单位宽度，全角约 2 个单位
        est = 0.0
        for seg in text.split("\n"):
            width_units = sum(2 if ord(ch) > 127 else 1 for ch in seg)
            est += max(1, math.ceil(width_units / max(4, w * 1.05)))
        lines = max(lines, int(est))
    ws.row_dimensions[row].height = max(min_h, min(max_h, lines * 12 + 4))


# ==================== 测试用例 ====================
def build_cases():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "测试用例"

    headers = ["用例编号", "用例标题", "项目/模块", "前置条件", "优先级",
               "测试步骤", "测试数据", "预期结果", "实际结果", "执行结论"]
    ncol = len(headers)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    t = ws.cell(row=1, column=1,
                value="ThinkRAG 本地知识库问答系统 — 测试用例与执行记录")
    t.font = TITLE_FONT
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    for i, h in enumerate(headers, start=1):
        ws.cell(row=2, column=i, value=h)
    style_header(ws, 2, ncol)

    row = 3
    widths = [15, 34, 12, 22, 7, 40, 26, 38, 48, 10]
    for c in CASES:
        values = [c["id"], c["title"], c["module"], c["pre"], c["pri"],
                  c["steps"], c["data"], c["expect"], c["actual"], c["result"]]
        for i, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = LEFT if i in (2, 4, 6, 7, 8, 9) else CENTER
        ws.cell(row=row, column=1).font = BOLD_BODY
        # 结果与优先级着色
        rc = ws.cell(row=row, column=ncol)
        rc.font = Font(name=FONT_NAME, size=9, bold=True,
                       color=RESULT_COLORS.get(c["result"], "000000"))
        pc = ws.cell(row=row, column=5)
        pc.font = Font(name=FONT_NAME, size=9, bold=True,
                       color=PRI_COLORS.get(c["pri"], "000000"))
        auto_row_height(ws, row, widths)
        row += 1

    set_widths(ws, [15, 34, 12, 22, 7, 40, 26, 38, 48, 10])
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(ncol)}{row - 1}"

    # ---------- 统计 ----------
    st = wb.create_sheet("统计")
    total = len(CASES)
    cnt = Counter(c["result"] for c in CASES)
    executed = cnt.get("通过", 0) + cnt.get("不通过", 0)

    st.cell(row=1, column=1, value="执行总览").font = TITLE_FONT
    st.merge_cells("A1:B1")
    st.cell(row=1, column=1).alignment = Alignment(horizontal="center")
    st.row_dimensions[1].height = 26

    overview = [
        ("指标", "数值"),
        ("用例总数", total),
        ("已执行", executed),
        ("通过", cnt.get("通过", 0)),
        ("不通过", cnt.get("不通过", 0)),
        ("待执行", cnt.get("待执行", 0)),
        ("阻塞", cnt.get("阻塞", 0)),
        ("通过率（按已执行计）", f"{cnt.get('通过', 0) / executed * 100:.1f}%" if executed else "-"),
        ("覆盖率（已执行/总数）", f"{executed / total * 100:.1f}%"),
    ]
    r = 2
    for k, v in overview:
        a = st.cell(row=r, column=1, value=k)
        b = st.cell(row=r, column=2, value=v)
        for cell in (a, b):
            cell.border = BORDER
            cell.font = BODY_FONT
            cell.alignment = CENTER
        if k == "指标":
            a.font = HEADER_FONT
            b.font = HEADER_FONT
            a.fill = HEADER_FILL
            b.fill = HEADER_FILL
        r += 1

    r += 1
    st.cell(row=r, column=1, value="自动化执行（pytest）").font = BOLD_BODY
    r += 1
    auto_rows = [
        ("自动化用例数", "39（另 1 条 slow 稳定性用例单独执行）"),
        ("通过", 27),
        ("已知缺陷（xfail）", 12),
        ("失败", 0),
        ("通过率（不含 xfail）", "100.0%"),
        ("CI 冒烟子集", "20 条 · MockEmbedding · 不依赖 LLM"),
        ("连续 10 次稳定性", "10/10 成功，P50 1.04s，P90 1.11s"),
        ("allure 报告", "reports/allure-report/index.html（39 条用例，含每次请求原文与耗时）"),
    ]
    for k, v in auto_rows:
        a = st.cell(row=r, column=1, value=k)
        b = st.cell(row=r, column=2, value=v)
        for cell in (a, b):
            cell.border = BORDER
            cell.font = BODY_FONT
            cell.alignment = CENTER
        a.fill = PatternFill("solid", fgColor="EAF3EA")
        r += 1

    r += 1
    st.cell(row=r, column=1, value="模块分布").font = BOLD_BODY
    r += 1
    for h, i in (("模块", 1), ("用例数", 2)):
        cell = st.cell(row=r, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = CENTER
    r += 1
    mod_counter = OrderedDict()
    for c in CASES:
        d = mod_counter.setdefault(c["module"], {"total": 0, "pass": 0, "fail": 0, "todo": 0})
        d["total"] += 1
        if c["result"] == "通过":
            d["pass"] += 1
        elif c["result"] == "不通过":
            d["fail"] += 1
        elif c["result"] == "待执行":
            d["todo"] += 1
    for mod, d in mod_counter.items():
        for i, v in enumerate([mod, f"{d['total']}（通过 {d['pass']} / 不通过 {d['fail']} / 待执行 {d['todo']}）"], start=1):
            cell = st.cell(row=r, column=i, value=v)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = CENTER if i == 1 else LEFT
        r += 1

    r += 1
    st.cell(row=r, column=1, value="优先级分布").font = BOLD_BODY
    r += 1
    for i, h in enumerate(["优先级", "数量"], start=1):
        cell = st.cell(row=r, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = CENTER
    r += 1
    pri_counter = Counter(c["pri"] for c in CASES)
    for p in ["P0", "P1", "P2", "P3"]:
        if pri_counter.get(p):
            for i, v in enumerate([p, pri_counter[p]], start=1):
                cell = st.cell(row=r, column=i, value=v)
                cell.font = BODY_FONT
                cell.border = BORDER
                cell.alignment = CENTER
            r += 1

    set_widths(st, [24, 46])

    # ---------- 实测指标 ----------
    mt = wb.create_sheet("实测指标")
    metrics = [
        ("指标", "数值", "采集方式"),
        ("服务健康检查", "HTTP 200", "curl /_stcore/health"),
        ("测试文档", "产品手册.txt，2743 字节 / 正文 1165 字", "文件属性"),
        ("chunk_size=2048 / overlap=512 切块数", "1 块", "统计向量库 embedding_dict 条数"),
        ("chunk_size=200 / overlap=50 切块数", "10 块", "向量库 10 条 / docstore 11 条"),
        ("chunk_size=100 / overlap=100", "53 块，总字符 5568（原文 1165，膨胀 4.78 倍）", "result_chunk.json"),
        ("chunk_size=512 / overlap=600", "10 块，总字符 5574（膨胀 4.78 倍）", "result_chunk.json"),
        ("chunk_size=1 / overlap=0", "ValueError: Metadata length (23) longer than chunk size (1)", "result_chunk.json"),
        ("格式矩阵可用", "txt / md / csv / docx / pdf / json 可解析", "result_format.json"),
        ("格式矩阵不可用", "pptx、xlsx 抛 ImportError；html 标签未剥离；exe、无扩展名按文本入库", "result_format.json"),
        ("空文件入库结果", "0 字节 → 1 个空节点", "result_format.json"),
        ("大文件切块耗时", "2.5MB → 2700 块，147.47 秒", "result_format.json"),
        ("重复上传", "docstore 11→22、向量 10→20（未去重）", "result_data.json"),
        ("同名文件覆盖", "向量 10→20，回答仍为旧值 1899 元", "result_data.json"),
        ("删除单文档", "docstore 14→11、向量 12→10，删除后正确拒答", "result_data.json"),
        ("多文档问答", "3/3 正确（企业采购 8 折、收纳包 129 元、保修 12 个月）", "result_data.json"),
        ("问答准确率", "事实题 11/11 命中", "result_qa.json"),
        ("超纲拒答率", "3/3 正确拒答（竞品对比、股价、单文档外事实）", "result_qa.json"),
        ("Top K 2 / 5 / 8 / 20", "found=2/5/8/11；TopK=2 时答案退化为拒答", "result_qa.json"),
        ("温度 0 一致性", "3 次逐字一致，不同答案数=1", "result_qa.json"),
        ("稳健性-超长输入", "1800 字问题 1.07 秒正常返回，参数全对", "result_robust.json"),
        ("稳健性-提示词注入", "正确拒绝，未泄露系统提示词", "result_robust.json"),
        ("稳健性-空字符串", "引擎层 TypeError（UI 层有守卫暂不可达）", "result_robust.json"),
        ("页面显示耗时", "0.03 / 0.02 / 0.02 秒", "页面 Took N second(s)"),
        ("引擎实测单次问答", "0.44 ~ 1.21 秒（14 次采样）", "result_qa.json"),
        ("嵌入模型单次实例化", "1.67 秒（首次）/ 0.17 秒（后续）", "create_embedding_model 连续 3 次计时"),
        ("索引规模（基线）", "1 个有效索引，10 个切块 + 1 条父文档", "index_store.json / docstore.json"),
        ("连续 10 次提问成功率", "100%（10/10，无异常、答案均非空）", "pytest -m slow / stability_10runs.json"),
        ("连续 10 次 P50 / P90", "1.04s / 1.11s（平均 0.95s，最大 1.12s）", "同上，排序后线性插值取位"),
        ("连续 10 次答案一致性", "4 种表述，核心事实完全一致（12 个月 / 6 个月）", "按答案原文去重比对"),
        ("pytest 自动化结果", "39 条：27 通过 / 12 xfail（已知缺陷）/ 0 失败", "pytest -m \"not slow\""),
        ("CI 冒烟子集", "20 条：10 通过 / 10 xfail，耗时约 14 秒", "GitHub Actions · MockEmbedding"),
    ]
    for ri, rowv in enumerate(metrics, start=1):
        for ci, v in enumerate(rowv, start=1):
            cell = mt.cell(row=ri, column=ci, value=v)
            if ri == 1:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
                cell.alignment = CENTER
            else:
                cell.font = BODY_FONT
                cell.alignment = LEFT if ci != 2 else CENTER
            cell.border = BORDER
    set_widths(mt, [34, 42, 40])

    path = os.path.join(BASE_DIR, "ThinkRAG测试用例.xlsx")
    wb.save(path)
    return path, total, executed, cnt


# ==================== 缺陷报告 ====================
def build_bugs():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "缺陷记录"

    headers = ["缺陷编号", "缺陷标题", "缺陷状态", "优先级", "模块", "缺陷描述", "附件"]
    ncol = len(headers)
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    style_header(ws, 1, ncol)

    row = 2
    widths_b = [14, 46, 10, 8, 12, 90, 30]
    for b in BUGS:
        values = [b["id"], b["title"], b["status"], b["pri"], b["module"], b["desc"], b["attach"]]
        for i, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = LEFT if i in (2, 6, 7) else CENTER
        ws.cell(row=row, column=1).font = BOLD_BODY
        sc = ws.cell(row=row, column=3)
        sc.font = Font(name=FONT_NAME, size=9, bold=True,
                       color=STATUS_COLORS.get(b["status"], "000000"))
        pc = ws.cell(row=row, column=4)
        pc.font = Font(name=FONT_NAME, size=9, bold=True,
                       color=PRI_COLORS.get(b["pri"], "000000"))
        auto_row_height(ws, row, widths_b, max_h=420)
        row += 1

    # 统计区
    stat_row = row + 1
    ws.cell(row=stat_row, column=1, value="统计").font = BOLD_BODY
    stat_row += 1
    cnt_status = Counter(b["status"] for b in BUGS)
    cnt_pri = Counter(b["pri"] for b in BUGS)
    cnt_mod = Counter(b["module"] for b in BUGS)
    lines = [
        ("缺陷总数", len(BUGS)),
        ("按状态", " / ".join(f"{k} {v}" for k, v in cnt_status.items())),
        ("按优先级", " / ".join(f"{k} {cnt_pri[k]}" for k in ["P1", "P2", "P3"] if cnt_pri.get(k))),
        ("按模块", " / ".join(f"{k} {v}" for k, v in cnt_mod.items())),
    ]
    for k, v in lines:
        a = ws.cell(row=stat_row, column=1, value=k)
        a.font = BOLD_BODY
        a.border = BORDER
        a.alignment = CENTER
        ws.merge_cells(start_row=stat_row, start_column=2, end_row=stat_row, end_column=5)
        bcell = ws.cell(row=stat_row, column=2, value=v)
        bcell.font = BODY_FONT
        bcell.alignment = LEFT
        bcell.border = BORDER
        stat_row += 1

    set_widths(ws, [14, 46, 10, 8, 12, 90, 30])
    ws.freeze_panes = "A2"

    # ---------- 缺陷闭环记录 ----------
    cs = wb.create_sheet("缺陷闭环记录")
    cs.merge_cells("A1:C1")
    tt = cs.cell(row=1, column=1, value="BUG-TRAG-001 闭环记录（发现 → 定位 → 提交 → 处置 → 验证 → 关闭）")
    tt.font = TITLE_FONT
    tt.alignment = Alignment(horizontal="center", vertical="center")
    cs.row_dimensions[1].height = 26

    for i, h in enumerate(["环节", "内容", "时间"], start=1):
        cs.cell(row=2, column=i, value=h)
    style_header(cs, 2, 3)

    r = 3
    for c in CLOSURE:
        for i, v in enumerate([c["step"], c["content"], c["date"]], start=1):
            cell = cs.cell(row=r, column=i, value=v)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = LEFT if i == 2 else CENTER
        cs.cell(row=r, column=1).font = BOLD_BODY
        r += 1

    set_widths(cs, [12, 95, 14])

    path = os.path.join(BASE_DIR, "ThinkRAG缺陷报告.xlsx")
    wb.save(path)
    return path


if __name__ == "__main__":
    p1, total, executed, cnt = build_cases()
    p2 = build_bugs()
    print(f"[OK] {p1}")
    print(f"     用例总数 {total} / 已执行 {executed} / 通过 {cnt.get('通过',0)} / "
          f"不通过 {cnt.get('不通过',0)} / 待执行 {cnt.get('待执行',0)}")
    print(f"[OK] {p2}")
    print(f"     缺陷总数 {len(BUGS)}")
